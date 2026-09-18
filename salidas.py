"""
Salidas de datos: CSV y XLSX con el historico acumulado.

El CSV sale en formato argentino (delimitador punto y coma, coma decimal)
porque Excel en es-AR interpreta el punto como separador de miles y rompe
las columnas numericas si se usa el formato ingles.
"""

from __future__ import annotations

import csv
import logging
import os
import shutil
from pathlib import Path
from typing import Iterable

from formato import a_local, formato_ar, iso_fecha

log = logging.getLogger(__name__)

# (clave interna, encabezado, tipo) -- el orden define el de las columnas.
COLUMNAS: list[tuple[str, str, str]] = [
    ("clave", "Clave", "texto"),
    ("fuente", "Fuente", "texto"),
    ("fecha_inicio", "Fecha licitacion", "fecha"),
    ("hora_inicio", "Apertura", "texto"),
    ("hora_fin", "Cierre", "texto"),
    ("estado", "Estado", "texto"),
    ("categoria", "Categoria", "texto"),
    ("titulo", "Titulo", "texto"),
    ("emisor", "Emisor", "texto"),
    ("tipo", "Tipo (API)", "texto"),
    ("moneda", "Moneda titulo", "texto"),
    ("moneda_monto", "Moneda monto", "texto"),
    # La clase normalizada, que es sobre la que filtra el bot. Vacia = la
    # publicacion no dice la moneda (tipico de los avisos de la CNV).
    ("moneda_clase", "Moneda (clase)", "texto"),
    ("monto_licitar", "Monto a licitar", "numero"),
    ("monto_adjudicado", "Monto adjudicado", "numero"),
    ("ratio_adjudicado", "Adjudicado / licitado", "ratio"),
    ("tasa_corte", "Tasa/margen de corte (%)", "tasa"),
    ("desierta", "Desierta", "bool"),
    ("valor_corte", "Valor de corte (texto)", "texto"),
    ("variable_licitar", "Variable licitada", "texto"),
    ("sistema_adjudicacion", "Sistema adjudicacion", "texto"),
    ("ampliable_hasta", "Ampliable hasta", "texto"),
    ("fecha_liquidacion", "Liquidacion", "fecha"),
    ("fecha_vencimiento", "Vencimiento", "fecha"),
    ("plazo_dias", "Plazo (dias)", "int"),
    ("industria", "Industria", "texto"),
    ("colocador", "Colocadores", "texto"),
    ("liquidador", "Liquidador", "texto"),
    ("rueda", "Rueda", "texto"),
    ("modalidad", "Modalidad", "texto"),
    ("observaciones", "Observaciones", "texto"),
    ("enlace", "Enlace", "texto"),
]


def enriquecer(reg: dict) -> dict:
    """Agrega las columnas derivadas que solo existen en la salida."""
    fila = dict(reg)

    inicio = a_local(reg.get("fecha_inicio"))
    fin = a_local(reg.get("fecha_fin"))
    fila["hora_inicio"] = inicio.strftime("%H:%M") if inicio else ""
    fila["hora_fin"] = fin.strftime("%H:%M") if fin else ""

    licitado = reg.get("monto_licitar")
    adjudicado = reg.get("monto_adjudicado")
    # Mide la demanda: un ratio > 1 significa que se amplio sobre lo ofrecido.
    fila["ratio_adjudicado"] = (
        adjudicado / licitado if licitado and adjudicado else None
    )

    venc = reg.get("fecha_vencimiento")
    liq = reg.get("fecha_liquidacion") or reg.get("fecha_inicio")
    fila["plazo_dias"] = (venc - liq).days if venc and liq and venc > liq else None

    return fila


def _valor_csv(fila: dict, clave: str, tipo: str) -> str:
    valor = fila.get(clave)
    if tipo == "fecha":
        return iso_fecha(valor)
    if tipo == "bool":
        return "SI" if valor else ""
    if valor is None or valor == "":
        return ""
    if tipo == "numero":
        return formato_ar(valor, 2)
    if tipo == "ratio":
        return formato_ar(valor, 4)
    if tipo == "tasa":
        return formato_ar(valor, 4)
    if tipo == "int":
        return str(int(valor))
    return str(valor).replace("\n", " ").replace("\r", " ")


def escribir_csv(registros: Iterable[dict], ruta: Path) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    filas = [enriquecer(r) for r in registros]

    # utf-8-sig: sin el BOM, Excel en Windows muestra los acentos rotos.
    with ruta.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow([enc for _, enc, _ in COLUMNAS])
        for fila in filas:
            w.writerow([_valor_csv(fila, c, t) for c, _, t in COLUMNAS])

    log.info("CSV escrito: %s (%d filas).", ruta, len(filas))
    return ruta


def copiar_a(archivos: Iterable[Path], destinos: Iterable[str]) -> list[str]:
    """Copia las salidas a carpetas adicionales, tipicamente un share de red.

    Devuelve la lista de problemas. Que un share no responda NO puede tirar
    abajo la corrida: el aviso por mail vale mas que la copia, y la carpeta de
    red es justo lo que se cae cuando alguien reinicia el servidor.

    Se escribe primero a un temporal en el destino y despues se renombra, para
    que nadie abra un dashboard a medio copiar.
    """
    problemas: list[str] = []

    for destino in destinos:
        if not str(destino).strip():
            continue
        carpeta = Path(destino)
        try:
            carpeta.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            problemas.append(f"No pude acceder a {destino}: {exc.strerror or exc}")
            log.error("No pude acceder a %s: %s", destino, exc)
            continue

        for origen in archivos:
            origen = Path(origen)
            if not origen.exists():
                continue
            final = carpeta / origen.name
            tmp = carpeta / f".{origen.name}.tmp"
            try:
                shutil.copyfile(origen, tmp)
                os.replace(tmp, final)
                log.info("Copiado %s -> %s", origen.name, carpeta)
            except OSError as exc:
                problemas.append(
                    f"No pude copiar {origen.name} a {destino}: {exc.strerror or exc}"
                )
                log.error("No pude copiar %s a %s: %s", origen.name, destino, exc)
                try:
                    tmp.unlink(missing_ok=True)
                except OSError:
                    pass

    return problemas


def escribir_xlsx(registros: Iterable[dict], ruta: Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        log.warning("openpyxl no esta instalado: salteo el XLSX (el CSV se genero igual).")
        return None

    filas = [enriquecer(r) for r in registros]

    wb = Workbook()
    ws = wb.active
    ws.title = "Licitaciones"

    encabezado_fill = PatternFill("solid", fgColor="1F2937")
    encabezado_font = Font(color="FFFFFF", bold=True, size=10)

    for col, (_, enc, _) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=enc)
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(vertical="center", wrap_text=False)

    formatos = {
        "numero": '#,##0',
        "ratio": '0.00%',
        "tasa": '0.00"%"',
        "fecha": 'dd/mm/yyyy',
    }

    for i, fila in enumerate(filas, start=2):
        for col, (clave, _, tipo) in enumerate(COLUMNAS, start=1):
            valor = fila.get(clave)
            if tipo == "fecha":
                local = a_local(valor)
                # Excel no maneja datetimes con zona: se pasa naive.
                valor = local.replace(tzinfo=None) if local else None
            elif tipo == "bool":
                valor = "SI" if valor else ""
            celda = ws.cell(row=i, column=col, value=valor)
            if tipo in formatos:
                celda.number_format = formatos[tipo]

    if filas:
        ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(filas) + 1}"
        tabla = Table(displayName="Licitaciones", ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(tabla)

    anchos = {
        "clave": 14, "fuente": 9, "enlace": 30,
        "titulo": 46, "emisor": 32, "observaciones": 50, "colocador": 34,
        "valor_corte": 34, "tipo": 26, "categoria": 20, "variable_licitar": 24,
        "sistema_adjudicacion": 18, "monto_licitar": 16, "monto_adjudicado": 16,
    }
    for col, (clave, enc, _) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = anchos.get(
            clave, max(11, min(len(enc) + 3, 22))
        )

    ws.freeze_panes = "C2"

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    log.info("XLSX escrito: %s (%d filas).", ruta, len(filas))
    return ruta
